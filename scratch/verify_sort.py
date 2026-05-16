import openpyxl
import os

path = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
s = wb['MOTHERDAY_ILOVEYOUMOM']

print("--- TOP 5 ROWS ---")
for r in range(3, 8):
    stt = s.cell(row=r, column=1).value
    camp = s.cell(row=r, column=2).value
    status = s.cell(row=r, column=6).value
    note = s.cell(row=r, column=5).value
    print(f"Row {r}: STT={stt}, Status='{status}', Note='{note}'")

print("\n--- BOTTOM 5 ROWS ---")
for r in range(s.max_row - 4, s.max_row + 1):
    if r < 3: continue
    stt = s.cell(row=r, column=1).value
    camp = s.cell(row=r, column=2).value
    status = s.cell(row=r, column=6).value
    note = s.cell(row=r, column=5).value
    print(f"Row {r}: STT={stt}, Status='{status}', Note='{note}'")
