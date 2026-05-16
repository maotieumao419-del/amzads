import openpyxl

path = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

# 1. Check sort order for a SKU sheet (Keyword before Product Targeting)
s = wb['MOTHERDAY_ILOVEYOUMOM']
print("=== MOTHERDAY_ILOVEYOUMOM - First 10 rows ===")
for r in range(3, 13):
    target = s.cell(row=r, column=4).value
    loai   = s.cell(row=r, column=3).value
    status = s.cell(row=r, column=6).value
    note   = s.cell(row=r, column=5).value
    print(f"  Row {r-2}: Status={status}, Type={loai}, Target={target}, Note={str(note)[:30]}")

# 2. Check Listing and Portfolio ID column widths
print("\n=== Listing Sheet - Column Widths ===")
ls = wb['Listing']
for i in range(1, 8):
    from openpyxl.utils import get_column_letter
    col = get_column_letter(i)
    print(f"  Col {col}: width={ls.column_dimensions[col].width:.1f}")

print("\n=== Portfolio ID Sheet - Column Widths ===")
pid = wb['Portfolio ID']
for i in range(1, 6):
    from openpyxl.utils import get_column_letter
    col = get_column_letter(i)
    print(f"  Col {col}: width={pid.column_dimensions[col].width:.1f}")
