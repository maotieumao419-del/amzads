import openpyxl
import os

path = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
s = wb['MOTHERDAY_ILOVEYOUMOM']

found_dummy = False
for r in range(3, s.max_row + 1):
    camp = s.cell(row=r, column=2).value
    status = s.cell(row=r, column=6).value
    if camp == 'DUMMY_CAMP':
        found_dummy = True
        impressions = s.cell(row=r, column=7).value
        print(f"Row {r}: Campaign='{camp}', Status='{status}', Impressions='{impressions}'")

if not found_dummy:
    print("DUMMY_CAMP not found in output!")

# Summary
statuses = [s.cell(row=r, column=6).value for r in range(3, s.max_row + 1)]
print(f"Enable: {statuses.count('Enable')}")
print(f"No-active: {statuses.count('No-active')}")
