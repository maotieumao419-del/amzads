import openpyxl
import os

files = [
    r"c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater - Copy\data\input 4\PPC_Musemory_UPDATED.xlsx",
    r"c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx"
]

for f in files:
    print(f"\nFULL PATH: {f}")
    if not os.path.exists(f):
        print("  FILE NOT FOUND")
        continue
    wb = openpyxl.load_workbook(f, data_only=True)
    import sys
    sys.stdout.reconfigure(encoding='utf-8')
    # Check the first SKU sheet (exclude Listing and Portfolio ID)
    sku_sheet = next((s for s in wb.sheetnames if s not in ['Listing', 'Portfolio ID']), None)
    if not sku_sheet:
        print("  NO SKU SHEET FOUND")
        continue
    ws = wb[sku_sheet]
    print(f"  Sheet: {sku_sheet}")
    # Print Row 1 (first 20 cols)
    row1 = [str(ws.cell(row=1, column=c).value) for c in range(1, 21)]
    print(f"  Row 1: {row1}")
    # Print Row 2 (first 20 cols)
    row2 = [str(ws.cell(row=2, column=c).value) for c in range(1, 21)]
    print(f"  Row 2: {row2}")
