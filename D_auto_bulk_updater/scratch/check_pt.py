import pandas as pd
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

output_file = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\input 4\PPC_Musemory_UPDATED.xlsx'
wb = openpyxl.load_workbook(output_file, data_only=True)

for sheet in wb.sheetnames:
    if sheet == 'Listing': continue
    ws = wb[sheet]
    for r in range(3, ws.max_row + 1):
        entity_type = str(ws.cell(row=r, column=3).value)
        if 'Product Targeting' in entity_type or 'Đối mục tiêu' in entity_type:
            campaign = ws.cell(row=r, column=2).value
            target = ws.cell(row=r, column=4).value
            # Check latest block
            next_col = 7
            while ws.cell(row=1, column=next_col).value is not None:
                current_block = next_col
                next_col += 12
            
            metrics = [ws.cell(row=r, column=current_block+i).value for i in range(11)]
            if all(v is None or v == 0 for v in metrics):
                print(f"Sheet: {sheet}, Row: {r}, Campaign: {campaign}, Target: {target}, Metrics: {metrics}")
                break
