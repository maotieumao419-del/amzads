
import openpyxl
import os

report_path = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\reports\PPC_Musemory_UPDATED_UNMATCHED_3004-0405.xlsx'

if not os.path.exists(report_path):
    print(f"File not found: {report_path}")
else:
    wb = openpyxl.load_workbook(report_path, data_only=True)
    if 'Bulk-Only (Not in Internal)' in wb.sheetnames:
        ws = wb['Bulk-Only (Not in Internal)']
        print(f"New Report Sample (Total rows: {ws.max_row}):")
        # Headers: Campaign Name, Target, Match Type, Impressions, Clicks, Spend, Predicted SKU, Reason
        for row in range(2, min(10, ws.max_row + 1)):
            data = [cell.value for cell in ws[row]]
            print(f" - [{data[7]}] {data[0]} | {data[1]} | {data[6]}")
    else:
        print("Sheet 'Bulk-Only' not found")
