import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

new = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
ref = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater - Copy\data\input 4\PPC_Musemory_UPDATED.xlsx'

wb_new = openpyxl.load_workbook(new, data_only=True)
wb_ref = openpyxl.load_workbook(ref, data_only=True)

if 'Portfolio ID' in wb_new.sheetnames:
    ws_n = wb_new['Portfolio ID']
    ws_r = wb_ref['Portfolio ID']
    print("=== Portfolio ID Comparison ===")
    print(f"  NEW: {ws_n.max_row} rows x {ws_n.max_column} cols")
    print(f"  REF: {ws_r.max_row} rows x {ws_r.max_column} cols")
    for r in range(1, 4):
        print(f"  Row {r} NEW: {[ws_n.cell(row=r, column=c).value for c in range(1, 6)]}")
        print(f"  Row {r} REF: {[ws_r.cell(row=r, column=c).value for c in range(1, 6)]}")
else:
    print("Portfolio ID sheet missing in NEW")
