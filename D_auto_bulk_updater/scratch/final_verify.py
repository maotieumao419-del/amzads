import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

new = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
ref = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater - Copy\data\input 4\PPC_Musemory_UPDATED.xlsx'

wb_new = openpyxl.load_workbook(new, data_only=True)
wb_ref = openpyxl.load_workbook(ref, data_only=True)

print("=== FINAL VERIFICATION ===")

# 1. Listing
ws_n = wb_new['Listing']
ws_r = wb_ref['Listing']
print("\n--- Listing Sheet ---")
print(f"  Header [1]: NEW={repr(ws_n.cell(row=1, column=1).value)}, REF={repr(ws_r.cell(row=1, column=1).value)}")
print(f"  SKU [2,2]:  NEW={repr(ws_n.cell(row=2, column=2).value)}, REF={repr(ws_r.cell(row=2, column=2).value)}")

# 2. Portfolio ID
ws_n = wb_new['Portfolio ID']
ws_r = wb_ref['Portfolio ID']
print("\n--- Portfolio ID Sheet ---")
print(f"  NEW: {ws_n.max_row}x{ws_n.max_column}, REF: {ws_r.max_row}x{ws_r.max_column}")
print(f"  Row 1 NEW: {[ws_n.cell(row=1, column=c).value for c in range(1, 6)]}")
print(f"  Row 1 REF: {[ws_r.cell(row=1, column=c).value for c in range(1, 6)]}")

# 3. ONM_NURSE
ws_n = wb_new['ONM_NURSE']
ws_r = wb_ref['ONM_NURSE']
print("\n--- ONM_NURSE Sheet ---")
print(f"  NEW: {ws_n.max_row} rows, REF: {ws_r.max_row} rows")
if ws_n.max_row == ws_r.max_row:
    print("  Row counts MATCH.")
else:
    print("  Row counts MISMATCH.")

# Check all sheets count
print(f"\nSheet count: NEW={len(wb_new.sheetnames)}, REF={len(wb_ref.sheetnames)}")
