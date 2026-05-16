import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

new = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
ref = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater - Copy\data\input 4\PPC_Musemory_UPDATED.xlsx'

wb_new = openpyxl.load_workbook(new, data_only=True)
wb_ref = openpyxl.load_workbook(ref, data_only=True)

ws_n = wb_new['Listing']
ws_r = wb_ref['Listing']

print("=== Listing Sheet Comparison ===")
print(f"  NEW: {ws_n.max_row} rows")
print(f"  REF: {ws_r.max_row} rows")

print("\n  Header (Row 1):")
print(f"    NEW: {[ws_n.cell(row=1, column=c).value for c in range(1, 9)]}")
print(f"    REF: {[ws_r.cell(row=1, column=c).value for c in range(1, 9)]}")

print("\n  First Data Row (Row 2):")
val_new = ws_n.cell(row=2, column=2).value
val_ref = ws_r.cell(row=2, column=2).value
print(f"    NEW SKU: {repr(val_new)}")
print(f"    REF SKU: {repr(val_ref)}")

# Check sheet names count and differences
new_sheets = set(wb_new.sheetnames)
ref_sheets = set(wb_ref.sheetnames)

print(f"\nSheet count: NEW={len(new_sheets)}, REF={len(ref_sheets)}")
print(f"Missing in NEW: {sorted(list(ref_sheets - new_sheets))}")
print(f"Extra in NEW:   {sorted(list(new_sheets - ref_sheets))[:10]}...")
