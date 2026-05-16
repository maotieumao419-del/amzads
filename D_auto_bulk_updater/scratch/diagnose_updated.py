import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

new = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\final_xlsx\PPC_Musemory_UPDATED.xlsx'
ref = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater - Copy\data\input 4\PPC_Musemory_UPDATED.xlsx'

wb_new = openpyxl.load_workbook(new, data_only=True)
wb_ref = openpyxl.load_workbook(ref, data_only=True)

for sname in ['ONM_NURSE']:
    ws_n = wb_new[sname]
    ws_r = wb_ref[sname]
    print(f"=== {sname} ===")
    print(f"  NEW: {ws_n.max_row} rows x {ws_n.max_column} cols")
    print(f"  REF: {ws_r.max_row} rows x {ws_r.max_column} cols")
    print(f"\n  NEW rows:")
    for r in range(1, ws_n.max_row+1):
        row = [ws_n.cell(row=r, column=c).value for c in range(1, 19)]
        print(f"    {r}: {row}")
    print(f"\n  REF rows:")
    for r in range(1, ws_r.max_row+1):
        row = [ws_r.cell(row=r, column=c).value for c in range(1, 19)]
        print(f"    {r}: {row}")
