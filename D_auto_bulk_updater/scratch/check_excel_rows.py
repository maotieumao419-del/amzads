
import openpyxl
import os

report_path = r'c:\Users\nnh16\ads-trading-system\TEST\D_auto_bulk_updater\data\reports\PPC_Musemory_UPDATED_UNMATCHED_3004-0405.xlsx'

if not os.path.exists(report_path):
    print(f"File not found: {report_path}")
else:
    wb = openpyxl.load_workbook(report_path, data_only=True)
    if 'Bulk-Only (Not in Internal)' in wb.sheetnames:
        ws = wb['Bulk-Only (Not in Internal)']
        print(f"Sheet 'Bulk-Only' has {ws.max_row} rows (including header)")
        for row in range(2, min(5, ws.max_row + 1)):
            print(f"Row {row}: {[cell.value for cell in ws[row]]}")
    else:
        print("Sheet 'Bulk-Only' not found")
